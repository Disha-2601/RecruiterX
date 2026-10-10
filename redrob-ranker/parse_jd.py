import os
import sys
import re
import json
from docx import Document

def parse_docx(file_path):
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n".join(full_text)

def parse_pdf(file_path):
    full_text = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                full_text.append(t)
    except Exception as e:
        print(f"pypdf extraction warning: {e}")
    
    if not full_text:
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        full_text.append(t)
        except Exception as e2:
            print(f"pdfplumber extraction warning: {e2}")
            
    return "\n".join(full_text)

def extract_role_title(text):
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    for line in lines[:15]:
        clean = line.strip('•*-# ').strip()
        if re.search(r"\b(developer|engineer|designer|manager|architect|lead|analyst|intern|full-stack|fullstack|frontend|backend|data|ai|ml|qa|tester|consultant|specialist)\b", clean, re.IGNORECASE):
            if len(clean) <= 60:
                return clean
            match = re.search(r"([A-Za-z\s\-/]+(Developer|Engineer|Designer|Manager|Architect|Lead|Analyst|Intern|Specialist))", clean, re.IGNORECASE)
            if match:
                return match.group(1).strip()
    return lines[0][:50] if lines else "Target Role"

def parse_jd_text(text):
    text_clean = text.replace('\xa0', ' ')
    
    # 0. Extract Title
    title = extract_role_title(text_clean)
    print(f"Extracted Title: {title}")

    # 1. Extract years of experience
    min_yoe = 0
    max_yoe = 50
    yoe_range_match = re.search(r"(\d+)\s*[\-–—\s]*to[\-–—\s]*(\d+)\s*years?", text_clean, re.IGNORECASE) or \
                      re.search(r"(\d+)\s*[\-–—]\s*(\d+)\s*years?", text_clean, re.IGNORECASE)
    if yoe_range_match:
        min_yoe = int(yoe_range_match.group(1))
        max_yoe = int(yoe_range_match.group(2))
        print(f"Extracted YoE range: {min_yoe} - {max_yoe} years")
    elif re.search(r"\bfresher\b", text_clean, re.IGNORECASE):
        min_yoe = 0
        max_yoe = 2
        print(f"Extracted YoE range: Fresher (0-2 years)")
    else:
        min_yoe_match = re.search(r"(\d+)\s*\+\s*years?", text_clean, re.IGNORECASE)
        if min_yoe_match:
            min_yoe = int(min_yoe_match.group(1))
            max_yoe = min_yoe + 5
            print(f"Extracted minimum YoE: {min_yoe}+ years")
            
    # 2. Extract locations
    known_locations = [
        "Noida", "Delhi", "Gurgaon", "Gurugram", "Delhi NCR", "NCR", "Bangalore", "Bengaluru",
        "Hyderabad", "Pune", "Mumbai", "Chennai", "Kolkata", "Ahmedabad", "Jaipur", "Chandigarh",
        "Kochi", "Indore", "Uttar Pradesh", "Remote", "India"
    ]
    preferred_locations = []
    for loc in known_locations:
        if re.search(rf"\b{re.escape(loc)}\b", text_clean, re.IGNORECASE):
            if loc in ["Delhi NCR", "NCR"] and "Delhi" in preferred_locations:
                continue
            if loc == "Bengaluru" and "Bangalore" in preferred_locations:
                continue
            preferred_locations.append(loc)
    print(f"Extracted preferred locations: {preferred_locations}")
    
    # 3. Extract work mode
    work_mode = "onsite"
    if re.search(r"on-site\s*/\s*hybrid|hybrid\s*/\s*on-site", text_clean, re.IGNORECASE):
        work_mode = "hybrid"
    elif re.search(r"\b(hybrid|remote|flexible)\b", text_clean, re.IGNORECASE):
        mode_match = re.search(r"\b(hybrid|remote|flexible)\b", text_clean, re.IGNORECASE).group(1).lower()
        work_mode = mode_match
    print(f"Extracted work mode: {work_mode}")
    
    # 4. Extract notice period
    notice_period = 30 # default
    notice_match = re.search(r"(\d+)\s*[\-–—]?\s*days?\s*(?:notice|preferred|joining)?", text_clean, re.IGNORECASE) or \
                   re.search(r"(\d+)\s*[\-–—]?\s*day\s*notice", text_clean, re.IGNORECASE)
    if notice_match and int(notice_match.group(1)) in [0, 7, 15, 30, 60, 90]:
        notice_period = int(notice_match.group(1))
    elif re.search(r"\bimmediate\b", text_clean, re.IGNORECASE):
        notice_period = 0
    print(f"Extracted notice period: {notice_period} days")
    
    # 5. Extract salary band
    salary_band = {"min": 0.0, "max": 0.0}
    salary_match = re.search(r"(?:₹|\bCTC\b|\bSalary\b|\bstipend\b)?[^\n\d]*(\d+(?:\.\d+)?)\s*[\-–—to]+\s*(\d+(?:\.\d+)?)\s*(LPA|Lakhs|L\b)", text_clean, re.IGNORECASE) or \
                   re.search(r"(\d+(?:\.\d+)?)\s*[\-–—to]+\s*(\d+(?:\.\d+)?)\s*(LPA|Lakhs|L\b)", text_clean, re.IGNORECASE)
    if salary_match:
        salary_band["min"] = float(salary_match.group(1))
        salary_band["max"] = float(salary_match.group(2))
        print(f"Extracted salary band: {salary_band['min']} - {salary_band['max']} LPA")
    else:
        stipend_match = re.search(r"(\d{2,3}),?(\d{3})\s*[\-–—to]+\s*(\d{2,3}),?(\d{3})", text_clean)
        if stipend_match:
            st1 = float(stipend_match.group(1) + stipend_match.group(2)) * 12 / 100000
            st2 = float(stipend_match.group(3) + stipend_match.group(4)) * 12 / 100000
            salary_band["min"] = round(st1, 1)
            salary_band["max"] = round(st2, 1)
            print(f"Extracted salary band from stipend: {salary_band['min']} - {salary_band['max']} LPA equivalent")
        
    # 6. Skill Dictionary
    skill_dictionary = {
        # Frontend
        "HTML": 1.0, "HTML5": 1.0, "CSS": 1.0, "CSS3": 1.0,
        "JavaScript": 1.0, "JavaScript (ES6+)": 1.0, "ES6+": 0.9, "ES6": 0.9,
        "React": 1.0, "React.js": 1.0, "Next.js": 0.9, "Vue": 0.9, "Vue.js": 0.9, "Angular": 0.9,
        "TypeScript": 1.0, "Responsive Design": 0.9, "REST APIs": 0.9, "REST API": 0.9,
        "GraphQL": 0.8, "Tailwind": 0.8, "Tailwind CSS": 0.8, "Bootstrap": 0.8,
        "Git & GitHub": 0.8, "Git": 0.8, "GitHub": 0.8, "UI/UX": 0.8, "Redux": 0.8,
        "Webpack": 0.7, "Vite": 0.7, "SASS": 0.7, "Cross-browser compatibility": 0.8,
        "Accessibility": 0.7, "Wireframes": 0.7, "State Management": 0.8,
        # Backend / Systems
        "Node.js": 1.0, "Express": 0.9, "Python": 1.0, "Django": 0.9, "Flask": 0.9, "FastAPI": 0.9,
        "Java": 1.0, "Spring Boot": 0.9, "C++": 1.0, "C#": 1.0, ".NET": 0.9, "Go": 1.0, "Golang": 1.0,
        "Microservices": 0.9, "gRPC": 0.8,
        # Databases
        "SQL": 0.8, "PostgreSQL": 0.9, "MySQL": 0.8, "MongoDB": 0.9, "Redis": 0.9,
        # Cloud & DevOps
        "Docker": 0.9, "Kubernetes": 0.9, "AWS": 0.9, "Azure": 0.9, "GCP": 0.9, "CI/CD": 0.8,
        # AI / ML / Data
        "Machine Learning": 0.9, "NLP": 0.9, "PyTorch": 0.9, "TensorFlow": 0.9,
        "RAG": 0.9, "Vector DBs": 0.9, "LLM": 0.9, "Scikit-Learn": 0.8, "Pandas": 0.8,
        "Airflow": 0.7, "Spark": 0.7, "Data Engineering": 0.7, "Evaluation Frameworks": 0.8
    }

    nice_dictionary = {
        "Git & GitHub": 0.4, "GitHub": 0.4, "Git": 0.4,
        "Portfolio": 0.4, "UI/UX": 0.4, "Wireframes": 0.3,
        "Docker": 0.4, "CI/CD": 0.4, "Unit Testing": 0.4,
        "Open-Source": 0.3, "TypeScript": 0.4, "Next.js": 0.4,
        "Fine-tuning LLMs": 0.4, "LoRA": 0.4, "XGBoost": 0.4, "Distributed Systems": 0.3
    }
    
    required_skills = {}
    nice_to_have_skills = {}
    
    text_lower = text_clean.lower()

    for skill_name, weight in skill_dictionary.items():
        pattern = r"(?:\b|\b)" + re.escape(skill_name.lower()) + r"(?:\b|\b)"
        if re.search(pattern, text_lower):
            required_skills[skill_name] = weight
            
    for skill_name, weight in nice_dictionary.items():
        pattern = r"(?:\b|\b)" + re.escape(skill_name.lower()) + r"(?:\b|\b)"
        if re.search(pattern, text_lower):
            if skill_name not in required_skills:
                nice_to_have_skills[skill_name] = weight

    if not required_skills:
        req_section_match = re.search(r"(?:Required Skills|Skills|Qualifications|Technical Requirements|Requirements)[:\n](.*?)(?:\n\n|\n[A-Z]|\Z)", text_clean, re.DOTALL | re.IGNORECASE)
        if req_section_match:
            raw_skills = req_section_match.group(1)
            skill_items = [s.strip('•*-· ') for s in re.split(r"[\n·,;]", raw_skills) if s.strip('•*-· ')]
            for item in skill_items[:8]:
                if len(item) < 35:
                    required_skills[item] = 0.9

    if not required_skills:
        required_skills = {
            "JavaScript": 1.0,
            "HTML/CSS": 1.0,
            "React.js": 0.9,
            "REST APIs": 0.9,
            "Git": 0.8
        }
    if not nice_to_have_skills:
        nice_to_have_skills = {
            "TypeScript": 0.4,
            "UI/UX": 0.4,
            "Next.js": 0.3
        }

    jd_requirements = {
        "title": title,
        "required_skills": required_skills,
        "nice_to_have_skills": nice_to_have_skills,
        "min_years_experience": min_yoe,
        "max_years_experience": max_yoe,
        "salary_band_inr_lpa": salary_band,
        "work_mode": work_mode,
        "preferred_locations": preferred_locations,
        "notice_period_days": notice_period,
        "jd_text_for_embedding": text_clean
    }
    return jd_requirements

def main():
    if len(sys.argv) < 2:
        print("Usage: python parse_jd.py <path_to_jd>")
        sys.exit(1)
        
    jd_path = sys.argv[1]
    if not os.path.exists(jd_path):
        print(f"Error: {jd_path} does not exist.")
        sys.exit(1)
        
    print(f"Parsing job description: {jd_path}")
    
    # Detect file type
    if jd_path.endswith('.docx'):
        text = parse_docx(jd_path)
    elif jd_path.endswith('.pdf'):
        text = parse_pdf(jd_path)
    else:
        # Check header bytes for pdf
        with open(jd_path, 'rb') as f:
            header = f.read(5)
        if header.startswith(b'%PDF'):
            text = parse_pdf(jd_path)
        else:
            with open(jd_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
            
    jd_requirements = parse_jd_text(text)
    
    os.makedirs("cache", exist_ok=True)
    
    # Save requirements json
    out_path = os.path.join("cache", "jd_requirements.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(jd_requirements, f, indent=2)
        
    # Also save raw text to cache/temp_jd.txt for rank.py / embed_jd.py
    txt_out_path = os.path.join("cache", "temp_jd.txt")
    with open(txt_out_path, "w", encoding="utf-8") as f:
        f.write(text)
        
    print(f"Successfully wrote JD requirements to {out_path} and plain text to {txt_out_path}")

if __name__ == "__main__":
    main()
